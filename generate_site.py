#!/usr/bin/env python3
"""
Dart League Site Generator
---------------------------
Reads the league's Excel workbook (Schedule / Results / Standings tabs) and
writes a static website (plain HTML/CSS, no server needed) that you can host
anywhere (GitHub Pages, Netlify, a folder on your own webhost, etc.).

USAGE
    python generate_site.py path/to/Dart_League_Schedule.xlsx

    Optional flags:
        --out DIR          output folder (default: ./site)
        --league-name NAME  shown in the header (default: see LEAGUE_NAME below)
        --logo PATH         image file to use as the league logo

IMPORTANT: this script reads the CACHED values of the workbook's formulas,
not the formulas themselves. That means the workbook must have been opened
and saved in Excel (or run through the recalc step) at least once after your
last edit, or the site will show blank/old numbers.

Re-run this script any time you update the workbook (new scores entered,
schedule changed, etc.) and re-upload the contents of the output folder to
your host.
"""

import argparse
import shutil
import sys
from pathlib import Path

import openpyxl

# ---- Defaults you can hand-edit here instead of passing flags every time ----
LEAGUE_NAME = "Wednesday Night Dart League"

SCHEDULE_FIRST_ROW = 7
SCHEDULE_LAST_ROW = 21

RESULTS_FIRST_ROW = 7
RESULTS_LAST_ROW = 36

LEADERBOARD_HEADER_ROW = 15
LEADERBOARD_FIRST_ROW = 16
LEADERBOARD_LAST_ROW = 21


def read_workbook(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    sched_ws = wb["Schedule"]
    schedule = []
    for r in range(SCHEDULE_FIRST_ROW, SCHEDULE_LAST_ROW + 1):
        night = sched_ws.cell(row=r, column=1).value
        date = sched_ws.cell(row=r, column=2).value
        board1 = sched_ws.cell(row=r, column=3).value
        board2 = sched_ws.cell(row=r, column=4).value
        bye = sched_ws.cell(row=r, column=5).value
        if night is None:
            continue
        schedule.append({
            "night": night, "date": date, "board1": board1,
            "board2": board2, "bye": bye,
        })

    res_ws = wb["Results"]
    results_by_night = {}
    for r in range(RESULTS_FIRST_ROW, RESULTS_LAST_ROW + 1):
        night = res_ws.cell(row=r, column=1).value
        if night is None:
            continue
        date = res_ws.cell(row=r, column=2).value
        board = res_ws.cell(row=r, column=3).value
        t1 = res_ws.cell(row=r, column=4).value
        t2 = res_ws.cell(row=r, column=5).value
        t1g = res_ws.cell(row=r, column=6).value
        t2g = res_ws.cell(row=r, column=7).value
        winner = res_ws.cell(row=r, column=9).value
        t1pts = res_ws.cell(row=r, column=10).value
        t2pts = res_ws.cell(row=r, column=11).value
        entry = {
            "board": board, "team1": t1, "team2": t2,
            "t1_games": t1g, "t2_games": t2g, "winner": winner,
            "t1_points": t1pts, "t2_points": t2pts,
        }
        results_by_night.setdefault(night, {"date": date, "matches": []})
        results_by_night[night]["matches"].append(entry)

    stand_ws = wb["Standings"]
    leaderboard = []
    for r in range(LEADERBOARD_FIRST_ROW, LEADERBOARD_LAST_ROW + 1):
        rank = stand_ws.cell(row=r, column=1).value
        team = stand_ws.cell(row=r, column=2).value
        if team is None:
            continue
        leaderboard.append({
            "rank": rank,
            "team": team,
            "wins": stand_ws.cell(row=r, column=3).value,
            "losses": stand_ws.cell(row=r, column=4).value,
            "points": stand_ws.cell(row=r, column=5).value,
            "games_won": stand_ws.cell(row=r, column=6).value,
        })

    return schedule, results_by_night, leaderboard


def fmt_date_short(d):
    """MM/DD/YYYY, used for the Schedule page's Date column."""
    if d is None:
        return ""
    if hasattr(d, "strftime"):
        return d.strftime("%m/%d/%Y")
    return str(d)


def fmt_date(d):
    if d is None:
        return ""
    try:
        return d.strftime("%a, %b %-d, %Y") if hasattr(d, "strftime") else str(d)
    except ValueError:
        return d.strftime("%a, %b %d, %Y")


def find_next_night(schedule, results_by_night):
    """First night with no results entered yet (or fully unplayed matches)."""
    for night_row in schedule:
        n = night_row["night"]
        played = results_by_night.get(n)
        if not played:
            return night_row
        if any(m["winner"] in (None, "", "Needs tiebreaker") for m in played["matches"]):
            return night_row
    return None


NAV_ITEMS = [("index.html", "Home"), ("schedule.html", "Schedule"),
             ("results.html", "Results"), ("standings.html", "Standings")]


def render_header(active_page, league_name):
    links = []
    for href, label in NAV_ITEMS:
        cls = ' class="active"' if href == active_page else ""
        links.append(f'<a href="{href}"{cls}>{label}</a>')
    return f"""
<header class="site-header">
  <div class="header-inner">
    <a href="index.html" class="brand">
      <img src="assets/logo.png" alt="" class="brand-logo" onerror="this.style.display='none'">
      <span class="brand-name">{league_name}</span>
    </a>
    <nav class="site-nav">{''.join(links)}</nav>
  </div>
</header>
"""


def render_footer():
    return """
<footer class="site-footer">
  <p>Scores and schedule generated from the league scorekeeping sheet.</p>
</footer>
"""


def page_shell(title, active_page, league_name, body):
    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{title} — {league_name}</title>
<link rel="stylesheet" href="style.css">
<link rel="preconnect" href="https://fonts.googleapis.com">
<link href="https://fonts.googleapis.com/css2?family=Bitter:ital,wght@0,400;0,600;0,700;1,400&family=JetBrains+Mono:wght@400;500;700&display=swap" rel="stylesheet">
</head>
<body>
{render_header(active_page, league_name)}
<main class="page">
{body}
</main>
{render_footer()}
</body>
</html>
"""


def render_next_up(night_row, results_by_night):
    if night_row is None:
        return '<div class="panel next-up"><p class="muted">Season complete — check Standings for final results.</p></div>'
    n, date = night_row["night"], fmt_date(night_row["date"])
    return f"""
<div class="panel next-up">
  <p class="eyebrow">Night {n} &middot; {date}</p>
  <div class="matchup-row">
    <div class="matchup"><span class="board-label">Board 1</span>{night_row['board1']}</div>
    <div class="matchup"><span class="board-label">Board 2</span>{night_row['board2']}</div>
  </div>
  <p class="bye-line"><span class="board-label">Bye</span>{night_row['bye']}</p>
</div>
"""


def render_standings_preview(leaderboard):
    rows = ""
    for row in leaderboard[:3]:
        rows += f"""<tr>
          <td class="rank">{row['rank']}</td>
          <td>{row['team']}</td>
          <td class="num">{row['points']}</td>
        </tr>"""
    return f"""
<div class="panel">
  <h2>Standings snapshot</h2>
  <div class="table-scroll">
  <table class="ledger compact">
    <thead><tr><th>Rank</th><th>Team</th><th>Points</th></tr></thead>
    <tbody>{rows}</tbody>
  </table>
  </div>
  <p><a href="standings.html">See full standings</a></p>
</div>
"""


def build_index(schedule, results_by_night, leaderboard, league_name):
    next_night = find_next_night(schedule, results_by_night)
    body = render_next_up(next_night, results_by_night) + render_standings_preview(leaderboard)
    return page_shell("Home", "index.html", league_name, body)


def build_schedule(schedule, league_name):
    rows = ""
    for row in schedule:
        rows += f"""<tr>
          <td class="num">{row['night']}</td>
          <td>{fmt_date_short(row['date'])}</td>
          <td>{row['board1']}</td>
          <td>{row['board2']}</td>
          <td class="muted">{row['bye']}</td>
        </tr>"""
    body = f"""
<h1>Schedule</h1>
<div class="panel">
  <div class="table-scroll">
  <table class="ledger">
    <thead><tr><th>Night</th><th>Date</th><th>Board 1</th><th>Board 2</th><th>Bye</th></tr></thead>
    <tbody>{rows}</tbody>
  </table>
  </div>
</div>
"""
    return page_shell("Schedule", "schedule.html", league_name, body)


def render_score(m):
    if m["t1_games"] in (None, "") or m["t2_games"] in (None, ""):
        return '<span class="muted">Not yet played</span>'
    score = f"{m['t1_games']}\u2013{m['t2_games']}"
    if m["winner"] in (None, "", "Needs tiebreaker"):
        return f'{score} <span class="muted">(tiebreaker needed)</span>'
    return score


def build_results(schedule, results_by_night, league_name):
    sections = ""
    for row in schedule:
        n = row["night"]
        night_data = results_by_night.get(n)
        matches_html = ""
        if night_data:
            for m in night_data["matches"]:
                winner = m["winner"]
                t1_cls = "winner" if winner == m["team1"] else ""
                t2_cls = "winner" if winner == m["team2"] else ""
                matches_html += f"""<tr>
                  <td class="num">{m['board']}</td>
                  <td class="{t1_cls}">{m['team1']}</td>
                  <td class="{t2_cls}">{m['team2']}</td>
                  <td class="num">{render_score(m)}</td>
                  <td class="num">{m['t1_points'] if m['t1_points'] not in (None,'') else '&ndash;'} / {m['t2_points'] if m['t2_points'] not in (None,'') else '&ndash;'}</td>
                </tr>"""
        sections += f"""
<div class="panel">
  <h2>Night {n} &middot; {fmt_date(row['date'])}</h2>
  <div class="table-scroll">
  <table class="ledger">
    <thead><tr><th>Board</th><th>Team 1</th><th>Team 2</th><th>Score</th><th>Points</th></tr></thead>
    <tbody>{matches_html}</tbody>
  </table>
  </div>
</div>
"""
    body = "<h1>Results</h1>" + sections
    return page_shell("Results", "results.html", league_name, body)


def build_standings(leaderboard, league_name):
    rows = ""
    for row in leaderboard:
        first_cls = " class=\"first-place\"" if row["rank"] == 1 else ""
        rows += f"""<tr{first_cls}>
          <td class="rank">{row['rank']}</td>
          <td>{row['team']}</td>
          <td class="num">{row['wins']}</td>
          <td class="num">{row['losses']}</td>
          <td class="num">{row['points']}</td>
          <td class="num">{row['games_won']}</td>
        </tr>"""
    body = f"""
<h1>Standings</h1>
<div class="panel">
  <div class="table-scroll">
  <table class="ledger">
    <thead><tr><th>Rank</th><th>Team</th><th>Wins</th><th>Losses</th><th>Points</th><th>Games won</th></tr></thead>
    <tbody>{rows}</tbody>
  </table>
  </div>
  <p class="muted footnote">Ties broken first by match points, then by total individual games won.</p>
</div>
"""
    return page_shell("Standings", "standings.html", league_name, body)


def main():
    parser = argparse.ArgumentParser(description="Generate the dart league static site from the schedule workbook.")
    parser.add_argument("xlsx", help="Path to the league Excel workbook")
    parser.add_argument("--out", default="site", help="Output folder (default: ./site)")
    parser.add_argument("--league-name", default=LEAGUE_NAME, help="League name shown in the header")
    parser.add_argument("--logo", default=None, help="Path to a logo image to use (png/jpg/svg)")
    args = parser.parse_args()

    xlsx_path = Path(args.xlsx)
    if not xlsx_path.exists():
        sys.exit(f"Could not find workbook: {xlsx_path}")

    out_dir = Path(args.out)
    assets_dir = out_dir / "assets"
    assets_dir.mkdir(parents=True, exist_ok=True)

    schedule, results_by_night, leaderboard = read_workbook(xlsx_path)

    script_dir = Path(__file__).parent

    def safe_copy(src, dst):
        """Copy src to dst, unless they're already the same file on disk."""
        if not src.exists():
            return
        if src.resolve() == dst.resolve():
            return  # already in place -- nothing to do
        shutil.copy(src, dst)

    safe_copy(script_dir / "style.css", out_dir / "style.css")

    if args.logo:
        logo_src = Path(args.logo)
        if logo_src.exists():
            safe_copy(logo_src, assets_dir / ("logo" + logo_src.suffix))
            if logo_src.suffix.lower() != ".png":
                print(f"Note: your logo isn't a .png. Rename the copied file in {assets_dir} to 'logo.png' "
                      f"(or edit generate_site.py's <img> tag) so it displays.")
    else:
        safe_copy(script_dir / "assets" / "logo.png", assets_dir / "logo.png")

    (out_dir / "index.html").write_text(build_index(schedule, results_by_night, leaderboard, args.league_name))
    (out_dir / "schedule.html").write_text(build_schedule(schedule, args.league_name))
    (out_dir / "results.html").write_text(build_results(schedule, results_by_night, args.league_name))
    (out_dir / "standings.html").write_text(build_standings(leaderboard, args.league_name))

    print(f"Site written to {out_dir.resolve()}")
    print(f"Open {out_dir.resolve() / 'index.html'} in a browser to preview it.")


if __name__ == "__main__":
    main()
